#!/usr/bin/env python3
"""Acquire public energy snapshots and build static catalogs. Requires openpyxl>=3.1.
Run without arguments to rebuild from retained, checksum-checked inputs; --refresh
re-downloads mutable upstream snapshots and updates their retrieval metadata.
"""
import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
from pathlib import Path
import re
import urllib.request
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data' / 'energy'
RAW = OUT / 'raw'
JURISDICTIONS = ['US', 'PR', 'VI', 'GU', 'AS', 'MP']
URLS = {
    **{f'geonames-{c}.zip': f'https://download.geonames.org/export/zip/{c}.zip' for c in JURISDICTIONS},
    'iou_zipcodes_2021.csv': 'https://data.openei.org/files/5806/iou_zipcodes_2021.csv',
    'non_iou_zipcodes_2021.csv': 'https://data.openei.org/files/5806/non_iou_zipcodes_2021.csv',
    'sales_revenue.xlsx': 'https://www.eia.gov/electricity/data/eia861m/xls/sales_revenue.xlsx',
    'egrid2023_data_rev2.xlsx': 'https://www.epa.gov/system/files/documents/2025-06/egrid2023_data_rev2.xlsx',
    'power_profiler_zipcode_tool_v14.2.xlsx': 'https://www.epa.gov/system/files/documents/2025-06/power_profiler_zipcode_tool_v14.2.xlsx',
}
SOURCES = {
    'geonames': {'publisher': 'GeoNames', 'title': 'Public postal-code inventory', 'url': 'https://download.geonames.org/export/zip/', 'license': 'CC BY 4.0 (download README declaration)', 'evidenceKind': 'public postal inventory; approximate coordinates', 'vintage': 'retrieval snapshot; not USPS-authoritative'},
    'oedi2021': {'publisher': 'NREL / OpenEI', 'title': 'U.S. Electric Utility Companies and Rates: Look-up by Zipcode (2021)', 'url': 'https://data.openei.org/submissions/5806', 'license': 'CC BY 4.0', 'vintage': 2021, 'published': '2022-11', 'evidenceKind': 'likely ZIP utility associations; not verified street service'},
    'eia861m': {'publisher': 'US Energy Information Administration', 'title': 'EIA-861M monthly state-sector electricity prices', 'url': URLS['sales_revenue.xlsx'], 'license': 'US public domain', 'licenseUrl': 'https://www.eia.gov/about/copyrights_reuse.php', 'evidenceKind': 'state-sector realized average retail price; survey estimates, preliminary/final status per row', 'units': 'USD/kWh converted from cents/kWh'},
    'egrid2023': {'publisher': 'US Environmental Protection Agency', 'title': 'eGRID2023 rev2', 'url': URLS['egrid2023_data_rev2.xlsx'], 'license': 'US public domain', 'licenseUrl': 'https://www.epa.gov/egrid/frequent-questions-about-egrid', 'vintage': 2023, 'published': '2025-06-12', 'evidenceKind': 'annual regional generation resource mix and CO2 total-output emission rate', 'units': {'mix': 'percent (0..100)', 'originalCO2': 'lb/MWh', 'co2KgPerKWh': 'kg CO2/kWh'}, 'conversion': 'lb/MWh * 0.45359237 / 1000 = kg/kWh'},
    'epaZip2023': {'publisher': 'US Environmental Protection Agency', 'title': 'Power Profiler ZIP tool v14.2', 'url': URLS['power_profiler_zipcode_tool_v14.2.xlsx'], 'license': 'US public domain', 'vintage': 'eGRID2023_rev1 associations, published 2025-06', 'evidenceKind': 'possible ZIP subregions; all alternatives retained; embedded rev1 rates unused'},
    'urdb': {'publisher': 'OpenEI', 'title': 'Utility Rate Database', 'url': 'https://openei.org/wiki/Utility_Rate_Database', 'bulkUrl': 'https://openei.org/apps/USURDB/download/usurdb.json.gz', 'license': 'OpenEI CC0 unless otherwise noted; underlying tariff documents may differ', 'evidenceKind': 'external tariff research source only; schedules not ingested or evaluated', 'coverageStatus': 'not-ingested'},
}


def emit(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n', encoding='utf-8')


def acquire(refresh):
    RAW.mkdir(parents=True, exist_ok=True)
    manifest_path = RAW / 'acquisition.json'
    prior = {r['file']: r for r in json.loads(manifest_path.read_text())} if manifest_path.exists() else {}
    records = []
    for name, url in URLS.items():
        path = RAW / name
        if refresh or not path.exists():
            request = urllib.request.Request(url, headers={'User-Agent': 'ControlWindow-public-data-ingestion/1.0'})
            with urllib.request.urlopen(request, timeout=180) as response:
                payload = response.read()
                modified = response.headers.get('Last-Modified')
            path.write_bytes(payload)
            record = {'file': name, 'url': url, 'bytes': len(payload), 'sha256': hashlib.sha256(payload).hexdigest(), 'retrievedAt': dt.datetime.now(dt.timezone.utc).isoformat(), 'lastModified': modified}
        else:
            if name not in prior:
                raise ValueError(f'{name}: cached file has no acquisition provenance; use --refresh')
            record = prior[name]
            if hashlib.sha256(path.read_bytes()).hexdigest() != record['sha256']:
                raise ValueError(f'{name}: source checksum changed; refusing an untracked build')
        records.append(record)
    emit(manifest_path, records)
    return records


def number(value):
    return value if isinstance(value, (int, float)) and math.isfinite(value) else None


def zipcode(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        value = str(int(value))
    value = str(value).strip()
    if not re.fullmatch(r'\d{1,5}', value):
        raise ValueError(f'Unexpected source ZIP {value!r}')
    return value.zfill(5)


def build(refresh=False):
    acquisition = acquire(refresh)
    zips = {}
    def ensure(code):
        return zips.setdefault(code, [code, None, None, None, None, None, 'mapping-only', [], []])
    source_rows = {}
    for country in JURISDICTIONS:
        with zipfile.ZipFile(RAW / f'geonames-{country}.zip') as archive:
            rows = list(csv.reader(io.StringIO(archive.read(f'{country}.txt').decode('utf-8')), delimiter='\t'))
            source_rows[f'geonames-{country}'] = len(rows)
            for row in rows:
                code = zipcode(row[1])
                current = ensure(code)
                state = row[4] if country == 'US' else country
                if current[6] == 'mapping-only':
                    current[1:7] = [row[2], state, float(row[9]) if row[9] else None, float(row[10]) if row[10] else None, int(row[11]) if row[11] else None, 'geonames']
                elif row[2] not in current[1].split(' / '):
                    current[1] += ' / ' + row[2]
    utilities, utility_ids = [], {}
    association_count = 0
    for filename in ['iou_zipcodes_2021.csv', 'non_iou_zipcodes_2021.csv']:
        with (RAW / filename).open(encoding='utf-8-sig', newline='') as handle:
            rows = csv.DictReader(handle)
            required = {'zip', 'eiaid', 'utility_name', 'state', 'service_type', 'ownership'}
            if not required.issubset(rows.fieldnames):
                raise ValueError(f'{filename}: utility schema changed')
            count = 0
            for row in rows:
                count += 1
                code = zipcode(row['zip'])
                entry = ensure(code)
                key = (row['eiaid'], row['utility_name'], row['state'], row['service_type'], row['ownership'])
                if key not in utility_ids:
                    utility_ids[key] = len(utilities)
                    utilities.append({'id': row['eiaid'], 'name': row['utility_name'], 'state': row['state'], 'serviceType': row['service_type'], 'ownership': row['ownership'], 'source': 'oedi2021', 'vintage': 2021})
                index = utility_ids[key]
                if index not in entry[7]:
                    entry[7].append(index)
                    association_count += 1
            source_rows[filename] = count
    # Mapping-only rows may carry a state inferred from the source association,
    # never a fabricated coordinate or a nearest-provider assignment.
    for entry in zips.values():
        if entry[2] is None:
            states = {utilities[i]['state'] for i in entry[7]}
            if len(states) == 1:
                entry[2] = next(iter(states))
    prices = {}
    workbook = openpyxl.load_workbook(RAW / 'sales_revenue.xlsx', read_only=True, data_only=True)
    for sheet in ['Monthly-States', 'Monthly-Ter']:
        rows = workbook[sheet].iter_rows(values_only=True)
        groups, fields, units = next(rows), next(rows), next(rows)
        if units[:4] != ('Year', 'Month', 'State', 'Data Status'):
            raise ValueError('EIA monthly header changed')
        columns = {}
        for sector in ['residential', 'commercial', 'industrial']:
            start = next(i for i, value in enumerate(groups) if str(value).lower() == sector)
            if fields[start + 3] != 'Price' or units[start + 3] != 'Cents/kWh':
                raise ValueError('EIA price units changed')
            columns[sector] = start
        count = 0
        for row in rows:
            if not isinstance(row[0], (int, float)) or not isinstance(row[1], (int, float)):
                continue
            state = str(row[2]).strip()
            if state == 'US':
                continue
            period = f'{int(row[0]):04d}-{int(row[1]):02d}'
            for sector, start in columns.items():
                price, sales = number(row[start + 3]), number(row[start + 1])
                if price is None or price <= 0 or sales is None or sales <= 0:
                    continue
                prices.setdefault(state, {}).setdefault(sector, []).append([period, round(price / 100, 6), row[3]])
                count += 1
        source_rows[sheet] = count
    workbook.close()
    price_coverage = {}
    for state, sectors in prices.items():
        price_coverage[state] = {}
        for sector, rows in sectors.items():
            rows.sort(key=lambda row: row[0])
            if len({row[0] for row in rows}) != len(rows):
                raise ValueError(f'Duplicate monthly prices: {state}/{sector}')
            price_coverage[state][sector] = {'firstPeriod': rows[0][0], 'lastPeriod': rows[-1][0], 'periods': len(rows), 'basis': 'state-sector-average-proxy'}
    fuels = {'CL': 'Coal', 'OL': 'Oil', 'GS': 'Natural gas', 'NC': 'Nuclear', 'HY': 'Hydro', 'BM': 'Biomass', 'WI': 'Wind', 'SO': 'Solar', 'GT': 'Geothermal', 'OF': 'Other fossil', 'OP': 'Other / purchased fuel'}
    grid = {'states': {}, 'subregions': {}}
    workbook = openpyxl.load_workbook(RAW / 'egrid2023_data_rev2.xlsx', read_only=True, data_only=True)
    for sheet, prefix, region_key, target in [('ST23', 'ST', 'PSTATABB', 'states'), ('SRL23', 'SR', 'SUBRGN', 'subregions')]:
        rows = workbook[sheet].iter_rows(values_only=True)
        labels, keys = next(rows), next(rows)
        if 'lb/MWh' not in labels[keys.index(prefix + 'CO2RTA')]:
            raise ValueError('eGRID emission units changed')
        for values in rows:
            row = dict(zip(keys, values))
            if number(row.get('YEAR')) is None or not row.get(region_key):
                continue
            co2 = number(row[prefix + 'CO2RTA'])
            mix = [{'fuel': fuel, 'percent': round(row[prefix + code + 'PR'] * 100, 4)} for code, fuel in fuels.items() if number(row.get(prefix + code + 'PR')) is not None]
            region = row[region_key]
            grid[target][region] = {'year': int(row['YEAR']), 'region': region, 'name': row.get('SRNAME') or region, 'geography': 'eGRID-subregion' if target == 'subregions' else 'state-generation', 'mix': mix, 'mixUnit': 'percent', 'generationMWh': number(row.get(prefix + 'NGENAN')), 'co2KgPerKWh': round(co2 * .45359237 / 1000, 9) if co2 is not None else None, 'originalCO2LbPerMWh': co2, 'co2Unit': 'kg CO2/kWh', 'evidenceKind': 'annual-regional-total-output-CO2', 'source': 'egrid2023'}
    workbook.close()
    workbook = openpyxl.load_workbook(RAW / 'power_profiler_zipcode_tool_v14.2.xlsx', read_only=True, data_only=True)
    rows = workbook['Zip-subregion'].iter_rows(values_only=True)
    if next(rows)[:4] != ('zip', 'Subregion 1', 'Subregion 2', 'Subregion 3'):
        raise ValueError('EPA ZIP-subregion header changed')
    count = 0
    for row in rows:
        if row[0] is None:
            continue
        code = zipcode(row[0])
        entry = ensure(code)
        for region in row[1:4]:
            if region and region not in entry[8]:
                entry[8].append(region)
        count += 1
    source_rows['epaZipSubregions'] = count
    workbook.close()
    entries = sorted(zips.values(), key=lambda row: row[0])
    coverage = {
        'zipRecords': len(entries), 'geonamesZipRecords': sum(row[6] == 'geonames' for row in entries),
        'mappingOnlyZipRecords': sum(row[6] == 'mapping-only' for row in entries),
        'zipsWithCoordinates': sum(row[3] is not None and row[4] is not None for row in entries),
        'zipsWithUtilityCandidates': sum(bool(row[7]) for row in entries),
        'zipsWithMultipleUtilityCandidates': sum(len(row[7]) > 1 for row in entries),
        'zipsUnknownUtility': sum(not row[7] for row in entries),
        'utilityAssociationCount': association_count, 'utilityCandidateRecords': len(utilities),
        'distinctEiaUtilityIds': len({u['id'] for u in utilities}),
        'zipsWithPriceHistory': sum(row[2] in prices for row in entries),
        'priceObservations': sum(len(rows) for sectors in prices.values() for rows in sectors.values()),
        'priceGeographies': len(prices), 'pricesByState': price_coverage,
        'zipsWithGridSubregions': sum(bool(row[8]) for row in entries),
        'zipsWithMultipleGridSubregions': sum(len(row[8]) > 1 for row in entries),
        'zipsWithStateGrid': sum(row[2] in grid['states'] for row in entries),
        'gridStateRecords': len(grid['states']), 'gridSubregionRecords': len(grid['subregions']), 'gridYear': 2023,
        'tariffSchedules': 'not-ingested; coverage unknown', 'sourceRows': source_rows,
        'inventoryClaim': 'Public GeoNames US/PR/VI/GU/AS/MP inventory plus retained mapping-only source ZIPs, not a USPS-authoritative active ZIP list. Mapping-only ZIPs may be obsolete or source placeholders.',
    }
    emit(ROOT / 'data' / 'us-zips.json', {'schemaVersion': 1, 'columns': ['zip', 'city', 'state', 'latitude', 'longitude', 'coordinateAccuracy', 'inventoryStatus', 'utilityIndexes', 'gridRegions'], 'rows': entries})
    emit(OUT / 'utilities.json', utilities)
    emit(OUT / 'prices.json', {'schemaVersion': 1, 'columns': ['period', 'usdPerKWh', 'status'], 'source': 'eia861m', 'states': prices})
    emit(OUT / 'grid.json', grid)
    emit(OUT / 'coverage.json', coverage)
    artifacts = []
    for path in [ROOT / 'data' / 'us-zips.json', OUT / 'utilities.json', OUT / 'prices.json', OUT / 'grid.json', OUT / 'coverage.json']:
        payload = path.read_bytes()
        artifacts.append({'file': str(path.relative_to(ROOT)), 'bytes': len(payload), 'sha256': hashlib.sha256(payload).hexdigest()})
    emit(OUT / 'manifest.json', {'schemaVersion': 1, 'parserVersion': '1.0', 'snapshotDate': max(r['retrievedAt'] for r in acquisition)[:10], 'sources': SOURCES, 'acquisition': acquisition, 'artifacts': artifacts, 'limitations': ['2021 ZIP utility candidates are not current verified premise service.', 'Monthly state-sector prices are proxies, not utility tariffs; actual observation periods only.', 'eGRID2023 is annual generation context, not utility procurement or hourly marginal emissions.', 'All ZIP source strings retained; mapping-only records are not evidence of active postal validity.', 'GeoNames coordinates are approximate and not suitable for verifying a utility service territory.']})
    print(json.dumps({key: value for key, value in coverage.items() if key != 'pricesByState'}, indent=2))
    print(json.dumps({'staticBytes': sum(r['bytes'] for r in artifacts), 'rawBytes': sum(r['bytes'] for r in acquisition), 'artifacts': artifacts}, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--refresh', action='store_true', help='Replace retained source files with fresh upstream snapshots')
    build(parser.parse_args().refresh)
